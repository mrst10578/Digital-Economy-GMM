"""Predictive relevance of actual lagged LEVEL instrument candidates for Difference GMM.

Prespecified 47-country leave-one-country-out predictive diagnostic.
Holds fixed BOTH ORIGINAL dependent-variable definitions Growth and ln_Productivity,
six contemporaneous explanatory variables, lag(2 3) collapsed instrument candidates,
and country-year 2008--2022 observations with complete lag2 AND lag3.

For dependent lag: target D.L.y_t = y_(t-1)-y_(t-2).
For six other variables: target D.x_t = x_t-x_(t-1).
For each of seven dynamic/regressor variables: z_lag2=x_(t-2), z_lag3=x_(t-3).
Within-year centered, fitted by QR (modified Gram-Schmidt, reorthogonalized).

This is deliberately NOT xtabond2's precise collapsed moment matrix,
NOT a formal Kleibergen-Paap / Stock-Yogo weak-IV test, NOT proof of exogeneity,
and NOT sufficient to declare a GMM estimator causally identified.
Country holdout assesses out-of-country instrument predictability, which can
be affected by country heterogeneity and the small effective T.
No country exclusion from original GMM estimates. No model/lag selection.
Runs using Python stdlib + openpyxl only, independent of Stata/Numpy.
"""
from __future__ import annotations

from collections import defaultdict
from hashlib import sha256
import csv
import json
import math
from pathlib import Path
from statistics import fmean
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "data/raw/Balanced_Panel_Data.xlsx"
RAW_SHA = "474d5525a37ecfbb424c9dd25d832a9ffa6eed1942bd4f4900454204810d215b"
OUT = ROOT / "outputs/difference_gmm_instrument_geometry"
OUTCOMES = ("Growth", "ln_Productivity")
X = ("Internet", "Broadband", "ICT_Exports", "Unemployment", "Inflation", "RnD")
YEAR_FIRST = 2008
YEAR_LAST = 2022
FOLDS = 47
Z_COLS = 14
TARGETS = 7

def load_panel() -> dict[str, dict[int, dict[str, float]]]:
    assert sha256(RAW.read_bytes()).hexdigest() == RAW_SHA, "Historical workbook changed"
    wb = load_workbook(RAW, read_only=True, data_only=True)
    assert len(wb.sheetnames) == 1
    rows = wb.active.iter_rows(values_only=True)
    header = list(next(rows))
    expected = ["Country", "ISO3", "Year", "Growth", "Productivity", *X]
    assert header == expected, (header, expected)
    panel: dict[str, dict[int, dict[str,float]]] = defaultdict(dict)
    for record in rows:
        if not record or all(v is None for v in record):
            continue
        item = dict(zip(header, record))
        iso = str(item["ISO3"])
        year = int(item["Year"])
        assert year in range(2005, 2023) and len(iso) == 3
        assert year not in panel[iso] and float(item["Productivity"]) > 0
        v = {n: float(item[n]) for n in ("Growth", *X)}
        v["ln_Productivity"] = math.log(float(item["Productivity"]))
        assert all(math.isfinite(z) for z in v.values())
        panel[iso][year] = v
    assert len(panel) == 47
    assert sum(map(len, panel.values())) == 846
    assert all(set(yrs) == set(range(2005, 2023)) for yrs in panel.values())
    return dict(panel)

def make_rows(panel: dict, outcome: str) -> list[dict[str, Any]]:
    assert outcome in OUTCOMES
    series = (outcome, *X)
    rows = []
    for iso in sorted(panel):
        yr = panel[iso]
        for year in range(YEAR_FIRST, YEAR_LAST + 1):
            t = [yr[year-1][outcome]-yr[year-2][outcome]]
            t.extend(yr[year][x]-yr[year-1][x] for x in X)
            z = []
            for v in series:
                z.extend((yr[year-2][v], yr[year-3][v]))
            assert len(t) == TARGETS and len(z) == Z_COLS
            rows.append({"iso": iso, "year": year, "y": tuple(t), "z": tuple(z)})
    assert len(rows) == 705 and len({r["iso"] for r in rows}) == FOLDS
    return rows

def dot(a: list[float], b: list[float]) -> float:
    return sum(x*y for x,y in zip(a,b))

def fit(training: list[dict]) -> tuple[list[list[float]], list[list[float]], dict, list[float], list[float], list[float]]:
    # Year means use TRAIN countries only so holdout has no target leakage.
    dates = sorted({r["year"] for r in training})
    ym = {}
    for year in dates:
        sub=[r for r in training if r["year"] == year]
        assert len(sub) in (46,47), (year,len(sub))
        ym[year] = {
            "z": [fmean(r["z"][i] for r in sub) for i in range(Z_COLS)],
            "y": [fmean(r["y"][i] for r in sub) for i in range(TARGETS)],
        }
    cz = [
        [(r["z"][i]-ym[r["year"]]["z"][i]) for r in training]
        for i in range(Z_COLS)
    ]
    # Center+scale per-column, same scale to train/test.
    std = [math.sqrt(dot(col,col)/len(col)) for col in cz]
    assert all(s > 1e-12 for s in std), "Near-constant instrument candidate"
    a = [[v / std[j] for v in col] for j,col in enumerate(cz)]

    Q: list[list[float]]=[]
    R=[[0.0]*Z_COLS for _ in range(Z_COLS)]
    diag=[]
    # Twice re-orthogonalized MGS for reproducibility without BLAS.
    for j, original in enumerate(a):
        vec=original[:]
        for _ in range(2):
            for i,q in enumerate(Q):
                k=dot(q,vec)
                R[i][j]+=k
                vec=[v-k*qv for v,qv in zip(vec,q)]
        pivot=math.sqrt(dot(vec,vec))
        diag.append(pivot)
        if pivot < 1e-10*math.sqrt(len(training)):
            raise ValueError("Instrument matrix loses numerical rank at column "+str(j))
        R[j][j] = pivot
        Q.append([v / pivot for v in vec])
    assert len(Q) == 14
    # Coefficients to predict all 7 target differences.
    betas=[]
    r2=[]
    for target in range(TARGETS):
        y=[r["y"][target] - ym[r["year"]]["y"][target] for r in training]
        qty=[dot(q,y) for q in Q]
        coefficients=[0.0]*Z_COLS
        for i in range(Z_COLS-1,-1,-1):
            coefficients[i]=(qty[i]-sum(R[i][j]*coefficients[j] for j in range(i+1,Z_COLS)))/R[i][i]
        den=dot(y,y)
        fitted=sum(k*k for k in qty)
        assert den>0
        within_r2=max(-1.0, min(1.0, fitted/den))
        betas.append(coefficients)
        r2.append(within_r2)
    spread = max(diag)/min(diag)
    assert math.isfinite(spread) and spread>=1
    return betas, Q, ym, std, diag, r2

def main() -> None:
    panel=load_panel()
    OUT.mkdir(parents=True,exist_ok=True)
    results=[]
    diagnostics={}
    for outcome in OUTCOMES:
        rows=make_rows(panel,outcome)
        betas, Q, mean, std, diag, in_r2=fit(rows)
        cumulative_sse=[0.0]*TARGETS
        cumulative_baseline=[0.0]*TARGETS
        by_country={}
        for iso in sorted(panel):
            train=[r for r in rows if r["iso"]!=iso]
            test=[r for r in rows if r["iso"]==iso]
            assert len(train)==690 and len(test)==15
            bb,_,ym,scale,piv,fold_r2=fit(train)
            numer=[0.0]*TARGETS
            denom=[0.0]*TARGETS
            for row in test:
                zv=[(row["z"][j]-ym[row["year"]]["z"][j])/scale[j] for j in range(Z_COLS)]
                for i in range(TARGETS):
                    actual=row["y"][i]-ym[row["year"]]["y"][i]
                    prediction=dot(bb[i],zv)
                    numer[i]+=(actual-prediction)**2
                    denom[i]+=actual**2
            for i in range(TARGETS):
                cumulative_sse[i]+=numer[i]
                cumulative_baseline[i]+=denom[i]
            by_country[iso]=[1-numer[i]/denom[i] if denom[i]>1e-20 else None for i in range(TARGETS)]
        vars=(outcome,*X)
        for i,v in enumerate(vars):
            denom=cumulative_baseline[i]
            assert denom>0
            oos_r2=1-cumulative_sse[i]/denom
            assert math.isfinite(oos_r2)
            results.append({
                "outcome":outcome,
                "differenced_endogenous_regressor":f"D.L.{outcome}" if i==0 else f"D.{v}",
                "full_sample_complete_case_rows":705,
                "country_holdout_folds":47,
                "lagged_instrument_columns":14,
                "within_year_in_sample_R2":f"{in_r2[i]:.10g}",
                "leave_one_country_out_predictive_R2":f"{oos_r2:.10g}",
                "oos_baseline_SSE":f"{denom:.10g}",
                "oos_instrument_SSE":f"{cumulative_sse[i]:.10g}",
                "instrument_assumptions_validated":"NO",
                "formal_GMM_weak_IV_test":"NO",
            })
        diagnostics[outcome]={
            "outcome":outcome,
            "instrument_column_count":Z_COLS,
            "numerically_full_column_rank":True,
            "QR_diagonal_max_min_ratio_not_spectral_condition_number":max(diag)/min(diag),
            "within_year_sample_rows":len(rows),
            "heldout_country_folds":FOLDS,
            "country_model_predictive_R2_by_target":by_country,
        }
    keys=list(results[0])
    with (OUT/"complete_instrument_predictive_relevance.csv").open("w",newline="",encoding="utf-8") as fp:
        writer=csv.DictWriter(fp,fieldnames=keys)
        writer.writeheader()
        writer.writerows(results)
    output={
        "status":"PASS: original 47 country data, 14 instrument candidate columns per outcome and 47 untouched country holdouts",
        "raw_data_sha256":RAW_SHA,
        "execution":"pure Python openpyxl; not Stata",
        "outcomes":list(OUTCOMES),
        "original_outcome_definitions_preserved":True,
        "complete_case_years":list(range(YEAR_FIRST,YEAR_LAST+1)),
        "complete_case_rows":705,
        "lag2_and_lag3_candidates_only":True,
        "formal_weak_IV_test":False,
        "GMM_exogeneity_proven":False,
        "GMM_causal_identification_validated":False,
        "limitations":[
            "Predictive leave-country-out R2 from within-year centered Z and regressor differences is not the conditional rank/weak identification diagnostic of xtabond2.",
            "The candidate matrix is a pooled conceptual proxy, not the actual block-diagonal time-specific GMM moment matrix.",
            "A negative R2 indicates poorer out-of-country prediction than the training-sample year mean, not invalid overidentifying moments.",
            "A positive R2 cannot establish exogeneity. Unmodeled country-specific relationships and temporal dependence can matter.",
            "Only complete years 2008-2022 with both lag2 and lag3 available: 705 rows rather than 752 for original GMM estimation.",
        ],
        "diagnostics":diagnostics,
        "table":results,
    }
    (OUT/"latest.json").write_text(json.dumps(output,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print("PREDICTIVE_AUDIT_PASS 2 outcomes x 7 endogenous targets x 47 country holdouts")
    for r in results:
        print("PREDICTIVE_RESULT",r["outcome"],r["differenced_endogenous_regressor"],
              "in_r2="+r["within_year_in_sample_R2"],
              "oos_r2="+r["leave_one_country_out_predictive_R2"])
    print("CAUSAL_MODELS_VALIDATED=0")

if __name__=="__main__":
    main()
