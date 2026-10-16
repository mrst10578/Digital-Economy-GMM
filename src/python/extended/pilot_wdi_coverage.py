#!/usr/bin/env python3
"""Phase B only: pilot WDI 47-country x 2005-2022 coverage. NO statistical fitting.

Downloads WDI API observations and source metadata, preserves raw responses,
builds a normalized 47x18 (including explicit missing) source table, and reports
coverage. The original professor workbook/branch are READ ONLY.

Do not treat additional controls as automatically exogenous IVs. A source's
coverage does not imply validity for causal identification.
"""
from __future__ import annotations
import csv
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import time
import urllib.parse
import urllib.request

ROOT = Path(__file__).resolve().parents[3]
LOCK = ROOT / "config/extended/phase2_original_47_country_sample_lock.json"
SOURCE = ROOT / "data/raw/Balanced_Panel_Data.xlsx"
OUT = ROOT / "outputs/extended/wdi_pilot_2026_10_16"
YEARS = range(2005, 2023)
INDICATORS = {
    "TM.VAL.ICTG.ZS.UN": "ICT goods imports (% of total goods imports)",
    "TX.VAL.TECH.MF.ZS": "High-technology exports (% of manufactured exports)",
    "IP.PAT.RESD": "Patent applications, residents",
    "IP.PAT.NRES": "Patent applications, nonresidents",
    "BX.KLT.DINV.WD.GD.ZS": "Foreign direct investment, net inflows (% of GDP)",
}

def get_json(url: str):
    last = None
    for attempt in range(4):
        try:
            req = urllib.request.Request(url,headers={"User-Agent":"Digital-Economy-GMM-Phase2-Coverage-Audit/1.0","Accept":"application/json"})
            with urllib.request.urlopen(req, timeout=35) as response:
                raw = response.read()
                if response.status != 200:
                    raise RuntimeError(f"HTTP {response.status}")
            obj = json.loads(raw)
            if isinstance(obj,dict) and obj.get("message"):
                raise RuntimeError(f"WDI API returned {obj['message']}")
            if not isinstance(obj,list) or len(obj)<2 or not isinstance(obj[1],list):
                raise RuntimeError(f"Unexpected API response from {url}, structure={str(obj)[:180]}")
            return obj, raw
        except Exception as error:
            last = str(error)
            print(f"API_RETRY attempt={attempt+1} error={last}",flush=True)
            if attempt < 3:
                time.sleep(1.4*(attempt+1))
    raise RuntimeError(f"WDI API unavailable; no missingness can be asserted. URL={url}, last={last}")

def csv_write(path, rows, fields):
    path.parent.mkdir(parents=True,exist_ok=True)
    with path.open("w",encoding="utf-8",newline="") as fp:
        w=csv.DictWriter(fp, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)

def main():
    lock=json.loads(LOCK.read_text(encoding="utf-8"))
    sh=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    assert sh==lock["original_data_sha256"], "Original workbook changed"
    isos=sorted(lock["iso3_to_original_country_name"].keys())
    assert len(isos)==47 and len(set(isos))==47
    assert lock["country_count"]==47 and lock["original_row_count"]==846
    # Confirm iso/year and completeness directly against untouched xlsx.
    import openpyxl
    wb=openpyxl.load_workbook(SOURCE,read_only=True,data_only=True)
    sht=wb.active
    hdr=[c.value for c in next(sht.iter_rows(min_row=1,max_row=1))]
    ci=hdr.index("ISO3"); yi=hdr.index("Year")
    input_keys=[(str(row[ci]),int(row[yi])) for row in sht.iter_rows(min_row=2,values_only=True)]
    full={(iso,year) for iso in isos for year in YEARS}
    assert len(input_keys)==846 and set(input_keys)==full
    assert len(set(input_keys))==846
    OUT.mkdir(parents=True, exist_ok=True)
    raw_dir=OUT / "raw"
    raw_dir.mkdir(exist_ok=True)
    source_map={}
    normalized={}
    base="https://api.worldbank.org/v2/"
    countries=";".join(isos)
    for code,desc in INDICATORS.items():
        params=urllib.parse.urlencode({"date":"2005:2022","per_page":"20000","format":"json"})
        url=f"{base}country/{countries}/indicator/{code}?{params}"
        response,payload=get_json(url)
        meta=response[0]
        assert int(meta.get("pages",1))==1, f"Unexpected paginated result for {code}: {meta}"
        archive=raw_dir / f"{code}.json"
        archive.write_bytes(payload)
        dictionary={}
        for o in response[1]:
            iso=o.get("countryiso3code") or o.get("country",{}).get("id")
            year=int(o["date"])
            if iso not in isos or year not in YEARS:
                continue
            key=(iso,year)
            assert key not in dictionary,(code,key)
            value=o.get("value")
            if value is not None:
                value=float(value)
                assert value==value and abs(value)<1e30,(code,key,value)
            dictionary[key]=value
        # If absent from API response, it is missing rather than zero.
        values={k:dictionary.get(k) for k in full}
        normalized[code]=values
        nonmiss=sum(v is not None for v in values.values())
        n_countries=sum(any(values[(iso,y)] is not None for y in YEARS) for iso in isos)
        coverage={
            "indicator":code,"description":desc,"source":"World Bank WDI V2",
            "available_47x18_cells":nonmiss,"potential_47x18_cells":846,
            "coverage_percent":round(100*nonmiss/846,2),
            "countries_with_any_data":n_countries,
            "countries_with_all_18_years":sum(all(values[(iso,y)] is not None for y in YEARS) for iso in isos),
            "countries_with_at_least_12_years":sum(sum(values[(iso,y)] is not None for y in YEARS)>=12 for iso in isos),
            "available_country_years_2005":sum(values[(iso,2005)] is not None for iso in isos),
            "available_country_years_2022":sum(values[(iso,2022)] is not None for iso in isos),
            "missing_not_imputed":True,
            "does_not_validate_exogeneity":True,
        }
        source_map[code]={"requested_url":url,"source_label":desc,
           "download_utc":datetime.now(timezone.utc).isoformat(),
           "raw_sha256":hashlib.sha256(payload).hexdigest(),"data_api_rows":len(response[1]),
           "metadata":meta,"coverage":coverage}
        print(f"INDICATOR_COVERAGE {code} {nonmiss}/846 {coverage['coverage_percent']}pct "
              f"countries_full={coverage['countries_with_all_18_years']}",flush=True)
    summary=[source_map[code]["coverage"] for code in INDICATORS]
    csv_write(OUT/"coverage_summary.csv",summary,list(summary[0]))
    by_country=[]
    for iso in isos:
        for code in INDICATORS:
            vals=normalized[code]
            n=sum(vals[(iso,y)] is not None for y in YEARS)
            by_country.append({"ISO3":iso,"source":code,"present_years":n,"missing_years":18-n,
                "coverage_percent":round(100*n/18,2)})
    csv_write(OUT/"country_coverage.csv",by_country,list(by_country[0]))
    value_rows=[]
    for iso in isos:
        for year in YEARS:
            row={"ISO3":iso,"Year":year}
            row.update({code:normalized[code][iso,year] for code in INDICATORS})
            value_rows.append(row)
    csv_write(OUT/"wdi_pilot_original_country_years.csv",value_rows,["ISO3","Year",*INDICATORS])
    (OUT/"provenance.json").write_text(json.dumps({
        "phase":"EXTENDED_DATA_EXPERIMENTAL",
        "original_data_sha256":sh,"original_47_country_sample_preserved":True,
        "original_data_modified":False,
        "only_pilot_new_variables_no_model_run":True,
        "year_window":"2005-2022","country_count":47,"rows":846,
        "dataset_license_to_verify":"WDI individual indicators and upstream ownership; confirm before reuse",
        "sources":source_map
    },ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    assert len(value_rows)==846
    print("PILOT_COMPLETE 5 WDI indicators x 47 countries x 18 years; no imputation; no modeling; Phase A unchanged")

if __name__=="__main__":
    main()
