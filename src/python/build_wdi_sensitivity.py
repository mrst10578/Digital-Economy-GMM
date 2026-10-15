"""Capture WDI revision sensitivity without overwriting the source workbook."""
from __future__ import annotations
import csv,datetime as dt,hashlib,json,urllib.parse,urllib.request
from pathlib import Path
from collections import Counter
import openpyxl
ROOT=Path(__file__).resolve().parents[2]
SOURCE=ROOT/"data/raw/Balanced_Panel_Data.xlsx"
OUT=ROOT/"outputs/quality_review"
VARMAP={"ICT_Exports":"BX.GSR.CCIS.ZS","Inflation":"FP.CPI.TOTL.ZG"}
COLUMNS=["column","indicator","ISO3","Year","original","current_WDI","delta","relative_change","request_url"]
def read_wdi(code):
    query=urllib.parse.urlencode({"format":"json","date":"2005:2022","per_page":20000})
    url=f"https://api.worldbank.org/v2/country/all/indicator/{code}?{query}"
    req=urllib.request.Request(url,headers={"User-Agent":"Research-Data-Provenance-Audit/2.0"})
    with urllib.request.urlopen(req,timeout=40) as resp: contents=json.load(resp)
    if not isinstance(contents,list) or len(contents)!=2 or int(contents[0].get("pages",1))!=1:
        raise ValueError("Invalid or paginated WDI response")
    return {(r["countryiso3code"],int(r["date"])):float(r["value"])
            for r in contents[1] if r and r.get("countryiso3code") and r.get("value") is not None},url
def main():
    OUT.mkdir(parents=True,exist_ok=True)
    sha=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    wb=openpyxl.load_workbook(SOURCE)
    ws=wb.active
    headings=[c.value for c in ws[1]]
    col={h:i+1 for i,h in enumerate(headings)}
    assert ws.max_row==847
    changes=[]
    for name,indicator in VARMAP.items():
        data,url=read_wdi(indicator)
        for line in range(2,ws.max_row+1):
            iso=ws.cell(line,col["ISO3"]).value
            year=int(ws.cell(line,col["Year"]).value)
            old=float(ws.cell(line,col[name]).value)
            v=data.get((iso,year))
            if v is None:raise ValueError(f"No current WDI data for {name} {iso} {year}")
            delta=v-old
            if abs(delta)<=max(1e-4,abs(v)*1e-5):continue
            relative=delta/max(abs(old),1e-8)
            changes.append({"column":name,"indicator":indicator,"ISO3":iso,"Year":year,
                            "original":old,"current_WDI":v,"delta":delta,
                            "relative_change":relative,"request_url":url})
            ws.cell(line,col[name]).value=v
        print("CURRENT_WDI_COL_REVISIONS",name,Counter(x["column"] for x in changes).get(name),flush=True)
    sha_after=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    assert sha_after==sha,"Source file was unexpectedly mutated"
    changes.sort(key=lambda x:(x["column"],x["ISO3"],x["Year"]))
    with (OUT/"wdi_revision_cells_20261015.csv").open("w",newline="",encoding="utf-8") as f:
        wr=csv.DictWriter(f,fieldnames=COLUMNS)
        wr.writeheader()
        wr.writerows(changes)
    notes=wb.create_sheet("WDI_CHANGE_AUDIT")
    notes.append(["Research sensitivity only, not an update to historical source",dt.date.today().isoformat()])
    notes.append(["Original SHA256",sha])
    notes.append(["Changed cells",len(changes)])
    notes.append(["column","ISO3","Year","original","current_WDI","delta"])
    for x in changes:
        notes.append([x["column"],x["ISO3"],x["Year"],x["original"],x["current_WDI"],x["delta"]])
    dest=OUT/"Balanced_Panel_Data_WDI_sensitivity_20261015.xlsx"
    wb.save(dest)
    count=Counter(x["column"] for x in changes)
    ranked=sorted(changes,key=lambda x:abs(x["delta"]),reverse=True)
    result={"source_sha256":sha,"sensitivity_file_sha256":hashlib.sha256(dest.read_bytes()).hexdigest(),
            "timestamp_utc":dt.datetime.now(dt.timezone.utc).isoformat(),
            "changed_cell_count":len(changes),"by_column":dict(count),"worst_absolute_revision":ranked[:12],
            "country_counts":dict(Counter(x["ISO3"] for x in changes)),
            "important":"Retains original in Git, revised workbook only for scenario comparison; not a replacement or proven correction."}
    (OUT/"wdi_revision_summary_20261015.json").write_text(json.dumps(result,indent=2,ensure_ascii=False))
    print("CURRENT_WDI_REVISIONS_SUMMARY",json.dumps(result,ensure_ascii=False)[:4800],flush=True)
if __name__=="__main__":main()
